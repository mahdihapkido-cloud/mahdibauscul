from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from pathlib import Path

class ExcelExporter:
    def __init__(self):
        Path("exports").mkdir(exist_ok=True)
    
    def export_records(self, records, filename=None):
        """
        صادرات رکوردها به Excel
        
        records: لیست رکوردها از دیتابیس
        filename: نام فایل (خودکار اگر نشود)
        """
        
        if filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"exports/weighing_records_{timestamp}.xlsx"
        
        wb = Workbook()
        ws = wb.active
        ws.title = "وزن‌گیری‌ها"
        
        # عنوان
        headers = [
            "شماره قبض",
            "پلاک",
            "وزن خالی (kg)",
            "وزن نهایی (kg)",
            "وزن خالص (kg)",
            "کد محصول",
            "تاریخ ثبت خالی",
            "تاریخ ثبت نهایی",
            "وضعیت"
        ]
        
        # استایل سرتیتر
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        # اضافه‌کردن هدرها
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        
        # استایل داده‌ها
        data_alignment = Alignment(horizontal="center", vertical="center")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # اضافه‌کردن رکوردها
        for row_num, record in enumerate(records, 2):
            _, _, empty_weight, empty_time, final_weight, final_time, product_code, \
            net_weight, receipt_number, status, _, _ = record
            
            row_data = [
                receipt_number or "-",
                record[1],  # plate_number
                empty_weight or "-",
                final_weight or "-",
                net_weight or "-",
                product_code or "-",
                empty_time or "-",
                final_time or "-",
                status or "-"
            ]
            
            for col_num, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col_num)
                cell.value = value
                cell.alignment = data_alignment
                cell.border = border
        
        # تنظیم عرض ستون‌ها
        column_widths = [15, 15, 15, 15, 15, 15, 20, 20, 12]
        for col_num, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(col_num)].width = width
        
        # تنظیم RTL برای فارسی
        ws.sheet_view.rightToLeft = True
        
        # ذخیره فایل
        wb.save(filename)
        return filename
    
    def export_daily_report(self, records, date_str=None):
        """صادرات گزارش روزانه"""
        
        if date_str is None:
            date_str = datetime.now().strftime("%Y-%m-%d")
        
        filename = f"exports/daily_report_{date_str}.xlsx"
        
        wb = Workbook()
        ws = wb.active
        ws.title = f"گزارش {date_str}"
        
        # عنوان
        ws['A1'] = f"گزارش روزانه وزن‌گیری - {date_str}"
        ws['A1'].font = Font(bold=True, size=14)
        ws.merge_cells('A1:I1')
        
        # خلاصه
        total_records = len(records)
        total_weight = sum([r[7] for r in records if r[7]])  # net_weight
        
        ws['A3'] = "خلاصه:"
        ws['A3'].font = Font(bold=True)
        ws['A4'] = f"تعداد وزن‌گیری: {total_records}"
        ws['A5'] = f"مجموع وزن خالص: {total_weight:.1f} کیلوگرم"
        
        # جدول
        headers = [
            "شماره قبض",
            "پلاک",
            "وزن خالی",
            "وزن نهایی",
            "وزن خالص",
            "کد محصول",
            "زمان",
        ]
        
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=7, column=col_num)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
        
        # اضافه‌کردن داده‌ها
        for row_num, record in enumerate(records, 8):
            ws.cell(row=row_num, column=1).value = record[8]  # receipt_number
            ws.cell(row=row_num, column=2).value = record[1]  # plate_number
            ws.cell(row=row_num, column=3).value = record[2]  # empty_weight
            ws.cell(row=row_num, column=4).value = record[4]  # final_weight
            ws.cell(row=row_num, column=5).value = record[7]  # net_weight
            ws.cell(row=row_num, column=6).value = record[6]  # product_code
            ws.cell(row=row_num, column=7).value = record[5]  # final_time
        
        # تنظیم عرض ستون‌ها
        column_widths = [12, 15, 12, 12, 12, 12, 18]
        for col_num, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(col_num)].width = width
        
        ws.sheet_view.rightToLeft = True
        
        wb.save(filename)
        return filename
